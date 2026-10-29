"""Actual income and expenses: the entry screen, the list, Excel import and budget against actual."""
import io
from decimal import Decimal

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..db import get_db
from ..models import ActualEntry, TdsSection, User
from ..security import can, current_user, deny, in_scope, require
from ..services import actuals as A
from ..services import budget as B

router = APIRouter(prefix="/api", tags=["actuals"])
viewer = require("view_actuals")
enterer = require("enter_actuals")


class EntryIn(BaseModel):
    date: str
    section: str
    head_code: str
    line_code: str
    amount: float
    party: str = ""
    voucher_type: str = ""
    voucher_no: str = ""
    narration: str = ""
    tds_code: str | None = None
    tds_amount: float = 0


def _check(db: Session, user: User, body: EntryIn):
    if body.section not in ("income", "expense"):
        raise HTTPException(422, "Choose income or expense.")
    if not can(user, f"view_{body.section}"):
        deny()
    fy, month = A.fy_of(body.date)
    year = B.year_by_start(db, fy)
    if not year:
        raise HTTPException(409, f"FY {fy}-{str(fy + 1)[2:]} is not set up. Add the year first.")
    hit = A.line_index(year).get((body.section, body.head_code, body.line_code))
    if not hit:
        raise HTTPException(422, "Choose a budget line from that year.")
    if not in_scope(user, hit[0]):
        deny("You can enter actuals only for your heads.")
    if not body.amount:
        raise HTTPException(422, "Enter the amount.")
    if body.tds_amount and body.section != "expense":
        raise HTTPException(422, "TDS applies to payments, not receipts.")
    if body.tds_amount and abs(body.tds_amount) > abs(body.amount):
        raise HTTPException(422, "TDS cannot be more than the amount.")
    if body.tds_code and not db.scalar(select(TdsSection).where(TdsSection.code == body.tds_code)):
        raise HTTPException(422, "That TDS category does not exist.")
    return fy, month, hit


def _apply(e: ActualEntry, body: EntryIn, fy: int, month: int):
    e.date, e.fy_start, e.month, e.section = body.date[:10], fy, month, body.section
    e.head_code, e.line_code = body.head_code, body.line_code
    e.amount, e.tds_amount = Decimal(str(round(body.amount, 2))), Decimal(str(round(body.tds_amount or 0, 2)))
    e.tds_code = body.tds_code or None
    e.party, e.voucher_type, e.voucher_no = body.party.strip()[:200], body.voucher_type.strip()[:60], body.voucher_no.strip()[:60]
    e.narration = body.narration.strip()[:500]


@router.get("/actuals")
def list_actuals(fy: int, section: str | None = None, month: int | None = None, head_code: str | None = None,
                 source: str | None = None, q: str | None = None, limit: int = Query(500, le=5000),
                 db: Session = Depends(get_db), user: User = Depends(viewer)):
    qy = select(ActualEntry).where(ActualEntry.fy_start == fy)
    if section:
        qy = qy.where(ActualEntry.section == section)
    if month is not None:
        qy = qy.where(ActualEntry.month == month)
    if head_code:
        qy = qy.where(ActualEntry.head_code == head_code)
    if source:
        qy = qy.where(ActualEntry.source == source)
    rows = [e for e in db.scalars(qy.order_by(ActualEntry.date.desc(), ActualEntry.id.desc()))
            if e.section == "unmapped" or can(user, f"view_{e.section}")]
    if q:
        ql = q.lower()
        rows = [e for e in rows if ql in f"{e.party} {e.narration} {e.ledger_name} {e.voucher_no}".lower()]
    year = B.year_by_start(db, fy)
    names = {}
    if year:
        for h in year.heads:
            for ln in h.lines:
                names[(h.code, ln.code)] = (h.name, ln.name)
    out = []
    for e in rows[:limit]:
        r = A.entry_out(e)
        r["head"], r["line"] = names.get((e.head_code, e.line_code), ("", ""))
        out.append(r)
    totals = {s: round(sum(float(e.amount) for e in rows if e.section == s), 2) for s in ("income", "expense")}
    return {"rows": out, "count": len(rows), "totals": totals}


@router.post("/actuals")
def add_actual(body: EntryIn, db: Session = Depends(get_db), user: User = Depends(enterer)):
    fy, month, (h, ln) = _check(db, user, body)
    e = ActualEntry(source="manual", created_by_name=user.name)
    _apply(e, body, fy, month)
    db.add(e)
    B.log(db, user, f"Actual {body.section}: ₹{body.amount:,.0f} on {ln.name} ({h.name}), {body.date}")
    db.commit()
    return A.entry_out(e)


@router.patch("/actuals/{entry_id}")
def edit_actual(entry_id: int, body: EntryIn, db: Session = Depends(get_db), user: User = Depends(enterer)):
    e = db.get(ActualEntry, entry_id)
    if not e:
        raise HTTPException(404, "That entry no longer exists.")
    if e.source == "tally":
        raise HTTPException(409, "Entries from Tally are changed in Tally. Change the voucher there and sync again.")
    fy, month, (h, ln) = _check(db, user, body)
    _apply(e, body, fy, month)
    B.log(db, user, f"Changed actual: ₹{body.amount:,.0f} on {ln.name} ({h.name}), {body.date}", key=f"actual:{e.id}")
    db.commit()
    return A.entry_out(e)


@router.delete("/actuals/{entry_id}")
def delete_actual(entry_id: int, db: Session = Depends(get_db), user: User = Depends(enterer)):
    e = db.get(ActualEntry, entry_id)
    if not e:
        raise HTTPException(404, "That entry no longer exists.")
    if e.source == "tally":
        raise HTTPException(409, "Entries from Tally are deleted in Tally. Delete the voucher there and sync again.")
    B.log(db, user, f"Deleted actual {e.section}: ₹{float(e.amount):,.0f}, {e.date}")
    db.delete(e)
    db.commit()
    return {"ok": True}


@router.get("/years/{year_id}/compare/{section}")
def compare(year_id: int, section: str, db: Session = Depends(get_db), user: User = Depends(viewer)):
    if section not in ("income", "expense"):
        raise HTTPException(404, "Unknown section.")
    if not can(user, f"view_{section}"):
        deny()
    return A.compare(db, B.get_year(db, year_id), section)


@router.get("/years/{year_id}/lines-index")
def lines_index(year_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    """Heads and lines of a year for pick lists (sections the user may see)."""
    y = B.get_year(db, year_id)
    out = {}
    for section in ("income", "expense"):
        if can(user, f"view_{section}"):
            out[section] = [{"code": h.code, "name": h.name, "can": in_scope(user, h),
                             "lines": [{"code": ln.code, "name": ln.name} for ln in h.lines]} for h in B.section_heads(y, section)]
    return out


# ---------------------------------------------------------------- Excel template and import
COLS = ["Date (YYYY-MM-DD)", "Income or expense", "Head", "Line", "Amount (₹)", "Party", "Voucher type", "Voucher no.", "Narration", "TDS category", "TDS deducted (₹)"]


@router.get("/years/{year_id}/actuals-template")
def template(year_id: int, db: Session = Depends(get_db), user: User = Depends(enterer)):
    y = B.get_year(db, year_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Actuals"
    ws.append(COLS)
    ws.append([f"{y.start_year}-04-10", "Expense", "Office Expenses", "Electricity", 18500, "MSEDCL", "Payment", "PV-101", "April bill", "", 0])
    ref = wb.create_sheet("Lines")
    ref.append(["Income or expense", "Head", "Line"])
    for h in y.heads:
        if can(user, f"view_{h.section}"):
            for ln in h.lines:
                ref.append(["Income" if h.section == "income" else "Expense", h.name, ln.name])
    codes = wb.create_sheet("TDS categories")
    codes.append(["Code", "Category", "Section (2025 Act)", "Old section"])
    for s in db.scalars(select(TdsSection).where(TdsSection.active.is_(True))):
        codes.append([s.code, s.label, s.new_ref, s.old_ref])
    dv = DataValidation(type="list", formula1='"Income,Expense"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B2:B2000")
    for c, w in zip("ABCDEFGHIJK", [16, 16, 28, 30, 14, 26, 14, 12, 30, 14, 16]):
        ws.column_dimensions[c].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="actuals-template-{y.label}.xlsx"'})


@router.post("/actuals/import")
async def import_actuals(file: UploadFile = File(...), dry_run: bool = Query(False), db: Session = Depends(get_db), user: User = Depends(enterer)):
    try:
        wb = load_workbook(io.BytesIO(await file.read()), data_only=True)
    except Exception:  # noqa: BLE001
        raise HTTPException(422, "That file could not be read. Use the actuals template.")
    ws = wb.worksheets[0]
    years: dict = {}
    good, errors = [], []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(v in (None, "") for v in row[:5]):
            continue
        row = list(row) + [None] * 11
        d = row[0]
        iso = d.date().isoformat() if hasattr(d, "date") else str(d or "").strip()[:10]
        section = str(row[1] or "").strip().lower()
        head, line = str(row[2] or "").strip().lower(), str(row[3] or "").strip().lower()
        try:
            fy, _m = A.fy_of(iso)
            if section not in ("income", "expense"):
                raise ValueError("write Income or Expense")
            if fy not in years:
                years[fy] = B.year_by_start(db, fy)
            y = years[fy]
            if not y:
                raise ValueError(f"FY {fy}-{str(fy + 1)[2:]} is not set up")
            hd = next((h for h in B.section_heads(y, section) if h.name.strip().lower() == head), None)
            ln = next((x for x in (hd.lines if hd else []) if x.name.strip().lower() == line), None)
            if not ln:
                raise ValueError(f"no line “{row[3]}” under head “{row[2]}”")
            body = EntryIn(date=iso, section=section, head_code=hd.code, line_code=ln.code, amount=float(row[4] or 0),
                           party=str(row[5] or ""), voucher_type=str(row[6] or ""), voucher_no=str(row[7] or ""),
                           narration=str(row[8] or ""), tds_code=(str(row[9]).strip().upper() or None) if row[9] else None,
                           tds_amount=float(row[10] or 0))
            _check(db, user, body)
            good.append(body)
        except HTTPException as e:
            errors.append(f"Row {i}: {e.detail}")
        except (ValueError, TypeError) as e:
            errors.append(f"Row {i}: {e}")
    if errors or dry_run:
        return {"ok": not errors, "rows": len(good), "errors": errors[:50]}
    for body in good:
        fy, month = A.fy_of(body.date)
        e = ActualEntry(source="import", created_by_name=user.name)
        _apply(e, body, fy, month)
        db.add(e)
    B.log(db, user, f"Imported {len(good)} actuals from {file.filename}")
    db.commit()
    return {"ok": True, "rows": len(good), "errors": []}
