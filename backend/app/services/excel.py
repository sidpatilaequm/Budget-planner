"""Excel export in the Budget_26-27.xlsx layout, and import from that layout."""
import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ..models import FiscalYear
from .budget import INSTALLMENTS, months, section_heads, spread

MONTH_NAMES = ["April", "May", "June", "July", "August", "September", "October", "November", "December",
               "January", "February", "March"]
MKEY = {"apr": 0, "may": 1, "jun": 2, "jul": 3, "aug": 4, "sep": 5, "oct": 6, "nov": 7, "dec": 8, "jan": 9,
        "feb": 10, "mar": 11}
SMALL = {"and", "of", "to", "in", "for", "the", "&"}
UPPER = {"xi", "xii", "x", "i", "ii", "rte", "npps", "dypps", "pcmc", "erp", "puc", "cbse", "ssc"}


def _clean(s) -> str:
    return re.sub(r"\s+", " ", str(s if s is not None else "")).strip()


def _title(s: str) -> str:
    words = _clean(s).lower().split(" ")
    out = []
    for i, w in enumerate(words):
        if w in UPPER:
            out.append(w.upper())
        elif i and w in SMALL:
            out.append(w)
        else:
            out.append(w[:1].upper() + w[1:])
    return " ".join(out)


def _isnum(v) -> bool:
    if v is None or v == "":
        return False
    try:
        float(v)
        return True
    except (TypeError, ValueError):
        return False


# ---------- export ----------

def _section_sheet(ws, year: FiscalYear, section: str):
    inc = section == "income"
    title = "Expenses" if section == "expense" else "Income"
    ws.append([f"{title} FY {year.label}"])
    ws["A1"].font = Font(bold=True, size=12)
    hdr = ["Sr No.", "Particulars"] + (["Qty", "Rate"] if inc else [])
    for k, m in enumerate(MONTH_NAMES):
        hdr.append(m)
        if k in (3, 7, 11):
            hdr.append(["1st", "2nd", "3rd"][k // 4] + " Installment")
    hdr += ["Total", "Percentage"]
    ws.append(hdr)
    for c in ws[2]:
        c.font = Font(bold=True)
    heads = section_heads(year, section)
    grand = sum(sum(months(l)) for h in heads for l in h.lines)

    def row(sr, name, q, r, m, bold=False):
        vals = [sr, name] + ([q if q is not None else "", r if r is not None else ""] if inc else [])
        for k, v in enumerate(m):
            vals.append(v)
            if k in (3, 7, 11):
                a, b = INSTALLMENTS[k // 4]
                vals.append(sum(m[a:b]))
        t = sum(m)
        vals += [t, (t / grand) if grand else 0]
        ws.append(vals)
        r_ = ws.max_row
        for ci in range(3, len(vals) + 1):
            cell = ws.cell(r_, ci)
            cell.number_format = "0.0%" if ci == len(vals) else "#,##,##0"
            if bold:
                cell.font = Font(bold=True)
        if bold:
            ws.cell(r_, 2).font = Font(bold=True)

    gm = [0.0] * 12
    for i, h in enumerate(heads, 1):
        ws.append([i, h.name])
        ws.cell(ws.max_row, 2).font = Font(bold=True)
        hm = [0.0] * 12
        hq = None
        for ln in h.lines:
            m = months(ln)
            row("", ln.name, float(ln.qty) if ln.qty is not None else None,
                float(ln.rate) if ln.rate is not None else None, m)
            hm = [a + b for a, b in zip(hm, m)]
            if ln.qty is not None:
                hq = (hq or 0) + float(ln.qty)
        row("", "Total", hq, None, hm, bold=True)
        gm = [a + b for a, b in zip(gm, hm)]
    row("", "TOTAL", None, None, gm, bold=True)
    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 44
    for ci in range(3, len(hdr) + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 13
    ws.freeze_panes = "C3"


def export_year(year: FiscalYear, sections: list[str]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for s in sections:
        name = ("Exps. " if s == "expense" else "Income ") + year.label[2:]
        _section_sheet(wb.create_sheet(name), year, s)
    if len(sections) == 2:
        ws = wb.create_sheet("Summary")
        ws.append([f"Summary FY {year.label}"])
        ws.append(["", "Income", "Expenses", "Net"])
        im = [sum(x) for x in zip(*[months(l) for h in section_heads(year, "income") for l in h.lines])] or [0] * 12
        em = [sum(x) for x in zip(*[months(l) for h in section_heads(year, "expense") for l in h.lines])] or [0] * 12
        for j, (a, b) in enumerate(INSTALLMENTS):
            i_, e_ = sum(im[a:b]), sum(em[a:b])
            ws.append([f"Inst. {['I', 'II', 'III'][j]}", i_, e_, i_ - e_])
        ws.append(["Year", sum(im), sum(em), sum(im) - sum(em)])
        for r in ws.iter_rows(min_row=3, min_col=2):
            for c in r:
                c.number_format = "#,##,##0"
        for col in "ABCD":
            ws.column_dimensions[col].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- import ----------

def _rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def parse_expense(ws) -> list[dict] | None:
    rows = _rows(ws)
    hi = next((i for i, r in enumerate(rows) if any(re.search("particulars", str(c or ""), re.I) for c in r)), None)
    if hi is None:
        return None
    hdr = rows[hi]
    name_col = next(i for i, c in enumerate(hdr) if re.search("particulars", str(c or ""), re.I))
    mcol = {}
    for ci, c in enumerate(hdr):
        k = str(c or "").strip()[:3].lower()
        if k in MKEY and MKEY[k] not in mcol:
            mcol[MKEY[k]] = ci
    heads, cur = [], None
    for r in rows[hi + 1:]:
        name = _clean(r[name_col] if name_col < len(r) else "")
        if not name or name.lower() == "total":
            continue
        if _isnum(r[0]):
            cur = {"name": _title(name), "lines": []}
            heads.append(cur)
            continue
        if cur is None:
            cur = {"name": "General", "lines": []}
            heads.append(cur)
        m = [float(r[mcol[k]]) if k in mcol and mcol[k] < len(r) and _isnum(r[mcol[k]]) else 0.0 for k in range(12)]
        cur["lines"].append({"name": re.sub(r"^[A-Z]\)\s*", "", name), "m": m, "qty": None, "rate": None})
    return heads or None


def parse_income(ws) -> list[dict] | None:
    rows = _rows(ws)
    heads, school, title, cur, qcol, rcol = [], "", "", None, -1, -1
    for r in rows:
        r = list(r) + [None] * 3
        a = _clean(r[0])
        filled = sum(1 for c in r if c is not None and _clean(c) != "")
        if re.match(r"^annex", a, re.I):
            title = _title(re.sub(r"^annex\s*-?\s*[a-z]\b", "", a, flags=re.I).strip() or "Income")
            cur = None
            continue
        if re.search("grand total", a, re.I):
            continue
        if filled == 1 and a and not _isnum(r[0]) and not re.match(r"^\d{4}", a):
            school = _title(a)
            continue
        b = _clean(r[1])
        if re.fullmatch(r"(std|class)", b, re.I):
            qcol = next((i for i, c in enumerate(r) if re.search("total new", str(c or ""), re.I)), -1)
            if qcol < 0:
                qcol = next((i for i, c in enumerate(r) if re.search("total student|strength", str(c or ""), re.I)), -1)
            rcol = next((i for i, c in enumerate(r) if i > 1 and re.search("fees", str(c or ""), re.I)
                         and not re.search("total", str(c or ""), re.I)), -1)
            cur = {"name": (school + " – " if school else "") + title, "lines": []}
            heads.append(cur)
            continue
        if cur is not None and _isnum(r[0]) and b:
            q = float(r[qcol]) if qcol >= 0 and _isnum(r[qcol]) else None
            rt = float(r[rcol]) if rcol >= 0 and _isnum(r[rcol]) else None
            m = spread(q * rt, "inst3") if q is not None and rt is not None else [0.0] * 12
            cur["lines"].append({"name": b, "m": m, "qty": q, "rate": rt})
    heads = [h for h in heads if h["lines"]]
    return heads or None


def parse_workbook(data: bytes) -> dict:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    out = {"start": None, "expense": None, "income": None}
    for ws in wb.worksheets:
        a1 = _clean(next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [None])[0])
        m = re.search(r"(20\d{2})\s*-\s*(\d{2,4})", a1)
        if m and not out["start"]:
            out["start"] = int(m.group(1))
        if re.search("exp", ws.title, re.I) and not out["expense"]:
            out["expense"] = parse_expense(ws)
        elif re.search("inc", ws.title, re.I) and not out["income"]:
            out["income"] = parse_income(ws)
    return out
