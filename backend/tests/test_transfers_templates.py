"""Budget transfers between lines, and the Excel template download and upload."""
import io
import os
import tempfile

if "DATABASE_URL" not in os.environ:
    os.environ["DATABASE_URL"] = os.getenv("TEST_DATABASE_URL", f"sqlite:///{tempfile.mkdtemp()}/t.db")

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.db import Base, engine  # noqa: E402
from app.main import app  # noqa: E402

APR = 0


@pytest.fixture(scope="module")
def client():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    with TestClient(app) as c:
        yield c


def H(t):
    return {"Authorization": f"Bearer {t}"}


@pytest.fixture(scope="module")
def ctx(client):
    r = client.post("/api/auth/setup", json={"name": "Ravi Kumar", "username": "ravi", "password": "secret123"})
    assert r.status_code == 200, r.text
    tok = r.json()["token"]
    year = client.get("/api/years", headers=H(tok)).json()[0]
    s = client.get(f"/api/years/{year['id']}/sections/expense", headers=H(tok)).json()
    heads = {h["name"]: h for h in s["heads"]}
    line = lambda head, name: next(l for l in heads[head]["lines"] if l["name"] == name)  # noqa: E731
    ids = {
        "maint": line("Transport Expenses", "Maintenance")["id"],
        "diesel": line("Transport Expenses", "Diesel Expenses")["id"],
        "insurance": line("Transport Expenses", "Bus Insurance")["id"],
        "elec": line("Office Expenses", "Electricity")["id"],
        "transport": heads["Transport Expenses"]["id"], "office": heads["Office Expenses"]["id"],
    }
    for key, amt in (("maint", 10000), ("diesel", 20000), ("elec", 8000)):
        r = client.patch(f"/api/lines/{ids[key]}", json={"m": [amt] + [0] * 11}, headers=H(tok))
        assert r.status_code == 200
    return {"tok": tok, "year": year, "ids": ids}


def month_value(client, ctx, line_id, month=APR):
    s = client.get(f"/api/years/{ctx['year']['id']}/sections/expense", headers=H(ctx["tok"])).json()
    return next(l["m"][month] for h in s["heads"] for l in h["lines"] if l["id"] == line_id)


def transfer(client, ctx, frm, to, amount, tok=None, month=APR, reason="Diesel prices fell"):
    return client.post(f"/api/years/{ctx['year']['id']}/transfers", headers=H(tok or ctx["tok"]),
                       json={"from_line_id": ctx["ids"][frm], "to_line_id": ctx["ids"][to], "month": month,
                             "amount": amount, "reason": reason})


# ------------------------------------------------------------------ transfers

def test_transfer_within_a_head_moves_budget_and_keeps_the_head_total(client, ctx):
    r = transfer(client, ctx, "diesel", "maint", 3000)
    assert r.status_code == 200, r.text
    assert r.json()["transfer"]["kind"] == "within"
    assert month_value(client, ctx, ctx["ids"]["diesel"]) == 17000
    assert month_value(client, ctx, ctx["ids"]["maint"]) == 13000
    assert client.get("/api/years", headers=H(ctx["tok"])).json()[0]["expense"] == 38000


def test_transfer_is_checked(client, ctx):
    assert transfer(client, ctx, "diesel", "maint", 17000.01).status_code == 422      # more than the line has
    assert transfer(client, ctx, "diesel", "maint", 100, month=5).status_code == 422  # nothing in September
    assert transfer(client, ctx, "diesel", "maint", 0).status_code == 422
    assert transfer(client, ctx, "diesel", "diesel", 10).status_code == 422
    assert transfer(client, ctx, "diesel", "maint", 10, reason=" ").status_code == 422
    assert month_value(client, ctx, ctx["ids"]["diesel"]) == 17000


def test_transfer_from_another_head_and_the_summary(client, ctx):
    r = transfer(client, ctx, "elec", "insurance", 2000, reason="Insurance renewal due")
    assert r.status_code == 200 and r.json()["transfer"]["kind"] == "cross"
    d = client.get(f"/api/years/{ctx['year']['id']}/transfers", headers=H(ctx["tok"])).json()
    assert len(d["transfers"]) == 2
    summary = {s["head"]: s for s in d["summary"]}
    assert summary["Office Expenses"]["out"] == 2000 and summary["Office Expenses"]["net"] == -2000
    assert summary["Transport Expenses"]["in"] == 2000 and summary["Transport Expenses"]["within"] == 3000
    assert client.get("/api/years", headers=H(ctx["tok"])).json()[0]["expense"] == 38000


def test_head_limit_blocks_budget_moving_in_but_not_within(client, ctx):
    y = ctx["year"]["id"]
    lim = client.put(f"/api/years/{y}/limits", headers=H(ctx["tok"]),
                     json={"heads": {str(ctx["ids"]["transport"]): {"mode": "amount", "value": 32000}}})
    assert lim.status_code == 200
    r = transfer(client, ctx, "elec", "insurance", 500)
    assert r.status_code == 409 and "over its limit" in r.json()["detail"]
    assert month_value(client, ctx, ctx["ids"]["elec"]) == 6000              # rolled back
    assert transfer(client, ctx, "maint", "diesel", 1000).status_code == 200  # within the head is fine
    client.put(f"/api/years/{y}/limits", headers=H(ctx["tok"]),
               json={"heads": {str(ctx["ids"]["transport"]): {"mode": "none"}}})


def test_reverse_a_transfer_once(client, ctx):
    d = client.get(f"/api/years/{ctx['year']['id']}/transfers", headers=H(ctx["tok"])).json()
    cross = next(t for t in d["transfers"] if t["kind"] == "cross")
    r = client.post(f"/api/transfers/{cross['id']}/reverse", json={}, headers=H(ctx["tok"]))
    assert r.status_code == 200 and r.json()["transfer"]["status"] == "reversal"
    assert month_value(client, ctx, ctx["ids"]["elec"]) == 8000
    assert client.post(f"/api/transfers/{cross['id']}/reverse", json={}, headers=H(ctx["tok"])).status_code == 409
    d = client.get(f"/api/years/{ctx['year']['id']}/transfers", headers=H(ctx["tok"])).json()
    assert {s["head"]: s["net"] for s in d["summary"]}["Office Expenses"] == 0


def _user(client, ctx, username, role_name, heads=None):
    roles = {r["name"]: r["id"] for r in client.get("/api/roles", headers=H(ctx["tok"])).json()}
    cat = client.get(f"/api/head-catalog?year_id={ctx['year']['id']}", headers=H(ctx["tok"])).json()
    body = {"name": username.title(), "username": username, "role_id": roles[role_name], "password": "pass12345",
            "scope": "some" if heads else "all", "heads": [h for h in cat if h["name"] in (heads or [])]}
    assert client.post("/api/users", json=body, headers=H(ctx["tok"])).status_code == 200
    return client.post("/api/auth/login", json={"username": username, "password": "pass12345"}).json()["token"]


def test_restricted_users_transfer_only_within_their_heads(client, ctx):
    dept = _user(client, ctx, "transport", "Department editor", ["Transport Expenses"])
    assert transfer(client, ctx, "maint", "diesel", 100, tok=dept).status_code == 200
    r = transfer(client, ctx, "elec", "maint", 100, tok=dept)
    assert r.status_code == 403 and "Office Expenses" in r.json()["detail"]
    viewer = _user(client, ctx, "viewer", "Viewer")
    assert transfer(client, ctx, "maint", "diesel", 100, tok=viewer).status_code == 403


# ------------------------------------------------------------------ templates

def _download(client, ctx, amounts=True, tok=None):
    r = client.get(f"/api/years/{ctx['year']['id']}/template?amounts={'true' if amounts else 'false'}",
                   headers=H(tok or ctx["tok"]))
    assert r.status_code == 200
    return load_workbook(io.BytesIO(r.content))


def _upload(client, ctx, wb, apply, tok=None):
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(f"/api/years/{ctx['year']['id']}/template/upload?apply={'true' if apply else 'false'}",
                    files={"file": ("t.xlsx", buf.getvalue())}, headers=H(tok or ctx["tok"]))
    assert r.status_code == 200, r.text
    return r.json()


def _row(ws, line_name):
    return next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 4).value == line_name)


def test_template_lists_every_head_and_line(client, ctx):
    wb = _download(client, ctx)
    assert wb.sheetnames[:3] == ["Instructions", "Expenses", "Income"]
    ws = wb["Expenses"]
    assert [c.value for c in ws[1]][:7] == ["Head code", "Line code", "Head", "Line", "Qty", "Rate", "April"]
    s = client.get(f"/api/years/{ctx['year']['id']}/sections/expense", headers=H(ctx["tok"])).json()
    assert ws.max_row - 1 == sum(len(h["lines"]) for h in s["heads"])
    assert ws.cell(_row(ws, "Electricity"), 7).value == 8000
    blank = _download(client, ctx, amounts=False)["Expenses"]
    assert blank.cell(_row(blank, "Electricity"), 7).value is None


def test_unchanged_template_changes_nothing(client, ctx):
    rep = _upload(client, ctx, _download(client, ctx), apply=True)
    assert rep["errors"] == [] and rep["applied"] is False
    assert rep["summary"]["updated"] == 0 and rep["summary"]["unchanged"] > 50


def test_upload_check_then_apply(client, ctx):
    wb = _download(client, ctx)
    ws = wb["Expenses"]
    r = _row(ws, "Electricity")
    ws.cell(r, 7).value = 9000          # April
    ws.cell(r, 8).value = 9000          # May
    ws.append([None, None, "Office Expenses", "Courier charges", None, None, 500, 500])   # new line, head by name
    ws.append([None, None, "Hostel", "Mess supplies", None, None, 25000])                # new head
    rep = _upload(client, ctx, wb, apply=False)
    assert rep["errors"] == [] and rep["applied"] is False
    assert rep["summary"] == {"updated": 1, "added_lines": 2, "added_heads": 1, "unchanged": rep["summary"]["unchanged"]}
    assert month_value(client, ctx, ctx["ids"]["elec"]) == 8000        # a check saves nothing
    rep = _upload(client, ctx, wb, apply=True)
    assert rep["applied"] is True
    assert month_value(client, ctx, ctx["ids"]["elec"]) == 9000
    s = client.get(f"/api/years/{ctx['year']['id']}/sections/expense", headers=H(ctx["tok"])).json()
    heads = {h["name"]: h for h in s["heads"]}
    assert heads["Hostel"]["lines"][0]["m"][0] == 25000
    assert any(l["name"] == "Courier charges" for l in heads["Office Expenses"]["lines"])
    # Uploading the same file again adds nothing twice.
    rep = _upload(client, ctx, wb, apply=True)
    assert rep["summary"]["added_lines"] == 0 and rep["summary"]["added_heads"] == 0 and rep["applied"] is False
    s = client.get(f"/api/years/{ctx['year']['id']}/sections/expense", headers=H(ctx["tok"])).json()
    names = [l["name"] for h in s["heads"] for l in h["lines"]]
    assert names.count("Courier charges") == 1 and names.count("Mess supplies") == 1


def test_a_bad_row_means_nothing_is_saved(client, ctx):
    wb = _download(client, ctx)
    ws = wb["Expenses"]
    ws.cell(_row(ws, "Electricity"), 7).value = 1
    ws.cell(_row(ws, "Diesel Expenses"), 9).value = "lots"
    ws.cell(_row(ws, "Maintenance"), 10).value = -5
    ws.append([None, "nosuchcode", "Office Expenses", "Ghost", None, None, 1])
    rep = _upload(client, ctx, wb, apply=True)
    assert rep["applied"] is False and len(rep["errors"]) == 3
    msgs = " ".join(e["message"] for e in rep["errors"])
    assert "not a number" in msgs and "negative" in msgs and "nosuchcode" in msgs
    assert month_value(client, ctx, ctx["ids"]["elec"]) == 9000


def test_restricted_user_upload_and_limits(client, ctx):
    dept = client.post("/api/auth/login", json={"username": "transport", "password": "pass12345"}).json()["token"]
    wb = _download(client, ctx, tok=dept)
    assert "Income" not in wb.sheetnames                       # cannot view income
    ws = wb["Expenses"]
    ws.cell(_row(ws, "Electricity"), 7).value = 1               # not their head
    ws.cell(_row(ws, "Road Tax"), 7).value = 4000               # their head
    rep = _upload(client, ctx, wb, apply=True, tok=dept)
    assert rep["applied"] is False and "Office Expenses" in rep["errors"][0]["message"]
    ws.cell(_row(ws, "Electricity"), 7).value = 9000            # back as downloaded
    assert _upload(client, ctx, wb, apply=True, tok=dept)["applied"] is True
    # A limit on the total blocks an upload that would exceed it.
    y = ctx["year"]["id"]
    client.put(f"/api/years/{y}/limits", headers=H(ctx["tok"]), json={"total": {"mode": "amount", "value": 1000}})
    wb = _download(client, ctx)
    ws = wb["Expenses"]
    ws.cell(_row(ws, "PUC"), 7).value = 700
    rep = _upload(client, ctx, wb, apply=True)
    assert rep["applied"] is False and "overall limit" in rep["errors"][0]["message"]
    client.put(f"/api/years/{y}/limits", headers=H(ctx["tok"]), json={"total": {"mode": "none"}})
    viewer = client.post("/api/auth/login", json={"username": "viewer", "password": "pass12345"}).json()["token"]
    buf = io.BytesIO(); wb.save(buf)
    r = client.post(f"/api/years/{y}/template/upload", files={"file": ("t.xlsx", buf.getvalue())}, headers=H(viewer))
    assert r.status_code == 403


def test_not_a_template(client, ctx):
    r = client.post(f"/api/years/{ctx['year']['id']}/template/upload", files={"file": ("x.xlsx", b"not excel")},
                    headers=H(ctx["tok"]))
    assert r.status_code == 422
    prog = client.get("/api/tools/upload_budget.py", headers=H(ctx["tok"]))
    assert prog.status_code == 200 and "def main" in prog.text
