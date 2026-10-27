"""The budget template: every head and line of a year in one workbook, filled in Excel and uploaded back.

Layout (one sheet per section, "Expenses" and "Income"):
    Head code | Line code | Head | Line | Qty | Rate | April … March | Total
The codes tie each row to a head and line in the planner. On upload:
- a row with a line code updates that line (amounts, quantity, rate and name);
- a row without a line code adds a new line under the head given by its head code or head name, unless that
  head already has a line of the same name, which is then updated (so uploading a file twice adds nothing twice);
- a head name that does not exist yet adds a new head (needs permission to manage heads);
- lines left out of the file are not touched, and nothing is deleted.
A row that cannot be used is reported with its sheet and row number. Changes are saved only when the
whole file is valid, so a file is either applied completely or not at all.
"""
import io
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import FiscalYear, User
from ..security import can, can_edit_amounts, can_edit_lines, can_manage_heads
from . import budget as B

SHEETS = {"expense": "Expenses", "income": "Income"}
MONTH_NAMES = ["April", "May", "June", "July", "August", "September", "October", "November", "December",
               "January", "February", "March"]
HEADER = ["Head code", "Line code", "Head", "Line", "Qty", "Rate"] + MONTH_NAMES + ["Total"]
COL = {name: i for i, name in enumerate(HEADER)}
FIRST_MONTH = COL["April"]
GREY = PatternFill("solid", fgColor="E7ECEF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E79")


# ---------------------------------------------------------------- download

def build_template(year: FiscalYear, sections: list[str], with_amounts: bool) -> bytes:
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    lines = [
        f"Budget template for FY {year.label}",
        "",
        "How to use this file",
        "1. Fill in the monthly amounts on the Expenses and Income sheets (April to March).",
        "2. To add a line (activity), add a row: leave Line code empty, give the Head (or Head code) and the Line name. "
        "If that head already has a line of that name, the row updates it instead, so uploading twice adds nothing twice.",
        "3. To add a head (group), add a row with a new Head name and its first Line. Head code stays empty.",
        "4. Do not change the Head code or Line code columns. They link each row to the planner.",
        "5. For a line priced by quantity, enter Qty and Rate and leave the months empty; the planner spreads it.",
        "6. Upload the file under Templates and upload. Check it first: nothing is saved until the whole file is valid.",
        "",
        "Rows you remove from the file are not deleted from the planner, and blank month cells count as zero.",
        "Amounts must not be negative. Only heads you may change can be changed; other rows must be left as they are.",
    ]
    for i, t in enumerate(lines, 1):
        ins.cell(i, 1, t)
    ins["A1"].font = Font(bold=True, size=14)
    ins["A3"].font = Font(bold=True)
    ins.column_dimensions["A"].width = 110

    for section in sections:
        ws = wb.create_sheet(SHEETS[section])
        ws.append(HEADER)
        for ci, _ in enumerate(HEADER, 1):
            c = ws.cell(1, ci)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEAD_FILL
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(1, 1).comment = Comment("Links the row to a head. Do not change it.", "Budget planner")
        ws.cell(1, 2).comment = Comment("Links the row to a line. Leave empty only for a new line.", "Budget planner")
        r = 2
        for h in B.section_heads(year, section):
            for ln in h.lines:
                m = B.months(ln) if with_amounts else [None] * 12
                row = [h.code, ln.code, h.name, ln.name,
                       float(ln.qty) if (with_amounts and ln.qty is not None) else None,
                       float(ln.rate) if (with_amounts and ln.rate is not None) else None] + \
                      [(v if v else None) for v in m]
                ws.append(row)
                first, last = get_column_letter(FIRST_MONTH + 1), get_column_letter(FIRST_MONTH + 12)
                ws.cell(r, len(HEADER), f"=SUM({first}{r}:{last}{r})")
                r += 1
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            row[0].fill = row[1].fill = GREY
            row[0].font = row[1].font = Font(color="5B6873", size=9)
            for c in row[4:]:
                c.number_format = "#,##,##0.00"
            row[-1].font = Font(bold=True)
        widths = [12, 12, 30, 38, 9, 11] + [12] * 12 + [14]
        for ci, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(ci)].width = w
        ws.freeze_panes = "E2"

    meta = wb.create_sheet("_meta")
    meta.append(["template", "budget-planner"])
    meta.append(["version", 1])
    meta.append(["year_start", year.start_year])
    meta.append(["generated", datetime.now().isoformat(timespec="seconds")])
    meta.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- upload

class TemplateError(Exception):
    pass


def _num(v):
    """Returns (number or None, error)."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, None
    if isinstance(v, (int, float)):
        return float(v), None
    s = str(v).strip().replace(",", "").replace("₹", "")
    try:
        return float(s), None
    except ValueError:
        return None, f"“{v}” is not a number"


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def read_template(data: bytes) -> dict:
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        raise TemplateError("That file could not be read as an Excel workbook (.xlsx).")
    out = {"year_start": None, "sections": {}}
    if "_meta" in wb.sheetnames:
        for row in wb["_meta"].iter_rows(values_only=True):
            if row and row[0] == "year_start" and isinstance(row[1], (int, float)):
                out["year_start"] = int(row[1])
    for section, title in SHEETS.items():
        if title not in wb.sheetnames:
            continue
        ws = wb[title]
        rows = list(ws.iter_rows(values_only=True))
        hdr_idx = next((i for i, r in enumerate(rows[:10]) if r and any(_text(c).lower() == "line code" for c in r)), None)
        if hdr_idx is None:
            raise TemplateError(f"The {title} sheet has no header row. Download a fresh template and use that.")
        hdr = [_text(c).lower() for c in rows[hdr_idx]]
        missing = [h for h in HEADER[:-1] if h.lower() not in hdr]
        if missing:
            raise TemplateError(f"The {title} sheet is missing the column(s): {', '.join(missing)}.")
        pos = {h: hdr.index(h.lower()) for h in HEADER[:-1]}
        parsed = []
        for n, r in enumerate(rows[hdr_idx + 1:], start=hdr_idx + 2):
            r = list(r) + [None] * (len(hdr) - len(r))
            get = lambda h: r[pos[h]]  # noqa: E731
            parsed.append({"row": n, "head_code": _text(get("Head code")), "line_code": _text(get("Line code")),
                           "head": _text(get("Head")), "line": _text(get("Line")),
                           "qty": get("Qty"), "rate": get("Rate"), "months": [get(m) for m in MONTH_NAMES]})
        out["sections"][section] = parsed
    if not out["sections"]:
        raise TemplateError("The file has no Expenses or Income sheet. Download a fresh template and use that.")
    return out


def process_upload(db, user: User, year: FiscalYear, parsed: dict, apply: bool) -> dict:
    """Checks every row and applies the changes in the session. The caller commits when `apply` is true and
    there are no errors, and rolls back otherwise."""
    errors, warnings, changes = [], [], []
    summary = {"updated": 0, "added_lines": 0, "added_heads": 0, "unchanged": 0}
    if parsed["year_start"] and parsed["year_start"] != year.start_year:
        warnings.append(f"This template was downloaded for FY {parsed['year_start']}-{str(parsed['year_start'] + 1)[2:]} "
                        f"and is being uploaded into FY {year.label}. Rows are matched by code, so this works for "
                        "heads and lines carried over between years.")
    before = B.snapshot_expense(year)

    for section, rows in parsed["sections"].items():
        sheet = SHEETS[section]
        if not can(user, f"view_{section}"):
            if any(r["line_code"] or r["line"] for r in rows):
                errors.append({"sheet": sheet, "row": None, "message": f"Your role cannot use the {sheet} sheet. Remove it from the file."})
            continue
        heads = B.section_heads(year, section)
        head_by_code = {h.code: h for h in heads}
        head_by_name = {h.name.strip().lower(): h for h in heads}
        line_by_code = {ln.code: ln for h in heads for ln in h.lines}
        next_order = max([h.sort_order for h in heads] or [-1]) + 1

        for r in rows:
            nums_present = any(v not in (None, "") for v in r["months"] + [r["qty"], r["rate"]])
            if not (r["head_code"] or r["line_code"] or r["head"] or r["line"] or nums_present):
                continue
            where = {"sheet": sheet, "row": r["row"]}
            row_errors = []
            months, bad = [], []
            for name, v in zip(MONTH_NAMES, r["months"]):
                x, e = _num(v)
                if e:
                    bad.append(f"{name}: {e}")
                elif x is not None and x < 0:
                    bad.append(f"{name} is negative")
                months.append(round(x or 0, 2))
            qty, eq = _num(r["qty"])
            rate, er = _num(r["rate"])
            bad += [f"Qty: {eq}"] if eq else []
            bad += [f"Rate: {er}"] if er else []
            if (qty is not None and qty < 0) or (rate is not None and rate < 0):
                bad.append("Qty and Rate must not be negative")
            if bad:
                row_errors.append("; ".join(bad))

            # A row without a line code whose line already exists under that head (by name) updates it,
            # so uploading the same file twice never adds the same line twice.
            matched = None
            if not r["line_code"] and r["line"]:
                h0 = (head_by_code.get(r["head_code"]) if r["head_code"]
                      else head_by_name.get(r["head"].lower()) if r["head"] else None)
                if h0:
                    matched = next((x for x in h0.lines if x.name.strip().lower() == r["line"].lower()), None)
            if r["line_code"] or matched:
                ln = matched or line_by_code.get(r["line_code"])
                if not ln:
                    row_errors.append(f"Line code {r['line_code']} is not in FY {year.label}. Clear the code to add this row as a new line.")
                elif r["head_code"] and r["head_code"] != ln.head.code:
                    row_errors.append(f"This line belongs to “{ln.head.name}”. Moving a line to another head is not supported; "
                                      "use a budget transfer instead.")
                if row_errors:
                    errors.append({**where, "message": " ".join(row_errors)})
                    continue
                h = ln.head
                if qty is not None and rate is not None and sum(months) == 0 and qty * rate > 0:
                    months = [float(v) for v in B.spread(qty * rate, ln.spread or "equal", ln.once_month, B.months(ln))]
                old_m = B.months(ln)
                old_q = float(ln.qty) if ln.qty is not None else None
                old_r = float(ln.rate) if ln.rate is not None else None
                amounts_changed = [round(a, 2) for a in old_m] != months or old_q != qty or old_r != rate
                name_changed = bool(r["line"]) and r["line"] != ln.name
                if not amounts_changed and not name_changed:
                    summary["unchanged"] += 1
                    continue
                if amounts_changed and not can_edit_amounts(user, h):
                    errors.append({**where, "message": f"You cannot change amounts in “{h.name}”. Leave this row as downloaded."})
                    continue
                if name_changed and not can_edit_lines(user, h):
                    errors.append({**where, "message": f"You cannot rename lines in “{h.name}”. Leave the Line name as downloaded."})
                    continue
                # Changes are made in the session either way; a check (or a file with errors) is rolled back.
                if amounts_changed:
                    B.set_months(ln, months)
                    ln.qty, ln.rate = qty, rate
                if name_changed:
                    ln.name = r["line"][:200]
                summary["updated"] += 1
                changes.append({**where, "action": "Updated" + (" and renamed" if name_changed and amounts_changed else
                                                              " name" if name_changed else ""),
                                "head": h.name, "line": ln.name, "before": sum(old_m), "after": sum(B.months(ln))})
                continue

            # ---- a new line
            if not r["line"]:
                row_errors.append("Enter the Line name for a new line, or the Line code of an existing line.")
            head = None
            if r["head_code"]:
                head = head_by_code.get(r["head_code"])
                if not head:
                    row_errors.append(f"Head code {r['head_code']} is not in FY {year.label}. Clear the code and give the Head name instead.")
            elif r["head"]:
                head = head_by_name.get(r["head"].lower())
            else:
                row_errors.append("Enter the Head for this line.")
            new_head = head is None and not r["head_code"] and bool(r["head"])
            if new_head and not can_manage_heads(user):
                row_errors.append(f"“{r['head']}” is not a head yet, and your role cannot add heads. Use an existing head name.")
            if head and not can_edit_lines(user, head):
                row_errors.append(f"You cannot add lines to “{head.name}”.")
            if new_head and can_manage_heads(user) and not (can(user, f"edit_{section}") and can(user, "edit_lines")):
                row_errors.append("Your role cannot add lines.")
            if row_errors:
                errors.append({**where, "message": " ".join(row_errors)})
                continue
            if new_head:
                head = B.add_head(db, year, section, r["head"][:160], order=next_order)
                next_order += 1
                head_by_name[head.name.lower()] = head
                summary["added_heads"] += 1
                changes.append({**where, "action": "Added head", "head": head.name, "line": "", "before": 0, "after": 0})
            ln = B.add_line(db, head, r["line"][:200], order=len(head.lines))
            if qty is not None and rate is not None and sum(months) == 0 and qty * rate > 0:
                months = [float(v) for v in B.spread(qty * rate, ln.spread, 0, None)]
            B.set_months(ln, months)
            ln.qty, ln.rate = qty, rate
            summary["added_lines"] += 1
            changes.append({**where, "action": "Added line", "head": head.name, "line": ln.name,
                            "before": 0, "after": sum(months)})

    if not errors:
        db.flush()
        msg = B.check_limits(db, year, before)
        if msg:
            if (year.limit_rule or "block") == "block":
                errors.append({"sheet": "Expenses", "row": None, "message": msg + ". Reduce other lines in the file, "
                               "or ask for the limit to be raised."})
            else:
                warnings.append(msg + ".")
    return {"summary": summary, "changes": changes, "errors": errors, "warnings": warnings}
